# -*- coding: utf-8 -*-
"""
Paso 2: extracción cruda de movimientos desde los 3 Excel (sin catálogos todavía).
Objetivo: dejar un CSV "crudo" con una fila por movimiento real, con el nombre de hoja/columna
tal cual viene en el archivo, para después mapearlo a los catálogos (equipos/categorías/proveedores).
No se descarta nada todavía; se valida el conteo total contra lo documentado en §9.4 (763/623/352).
"""
import openpyxl
import csv

SKIP_2024 = {"CONCENTRADO", "GASTOS DEPTO", "RENTAS"}
SKIP_2025 = {"CONCENTRADO"}

rows_out = []

def to_num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except Exception:
        return None

# ---------- 2024 ----------
wb = openpyxl.load_workbook("../privado/COSTOS POR UNIDAD 2024.xlsx", data_only=True, read_only=True)
for sheet in wb.sheetnames:
    if sheet in SKIP_2024:
        continue
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    header = [ (h or "").strip() if isinstance(h,str) else h for h in header ]
    idx = {h:i for i,h in enumerate(header) if h}
    n = 0
    for row in it:
        dia = row[idx.get("DIA", -1)] if idx.get("DIA",-1)>=0 else None
        if dia is None:
            continue
        n += 1
        rows_out.append({
            "archivo_origen":"2024","hoja_origen":sheet,
            "dia": row[idx["DIA"]], "mes": row[idx["MES"]], "anio": 2024,
            "factura": row[idx.get("FACTURA",-1)] if idx.get("FACTURA",-1)>=0 else None,
            "proveedor": row[idx.get("PROVEEDOR",-1)] if idx.get("PROVEEDOR",-1)>=0 else None,
            "concepto": row[idx.get("CONCEPTO",-1)] if idx.get("CONCEPTO",-1)>=0 else None,
            "equipo_col": None, "empleado_col": None,
            "cantidad": row[idx.get("CANTIDAD",-1)] if idx.get("CANTIDAD",-1)>=0 else None,
            "unidad": row[idx.get("UNIDAD",-1)] if idx.get("UNIDAD",-1)>=0 else None,
            "importe": to_num(row[idx.get("IMPORTE ",idx.get("IMPORTE",-1))]) if (idx.get("IMPORTE ",idx.get("IMPORTE",-1)))>=0 else None,
            "iva": to_num(row[idx.get("IVA",-1)]) if idx.get("IVA",-1)>=0 else None,
            "bruto": None,
            "neto": to_num(row[idx.get("NETO",-1)]) if idx.get("NETO",-1)>=0 else None,
            "tipo": row[idx.get("TIPO",-1)] if idx.get("TIPO",-1)>=0 else None,
        })
    print("2024", sheet, "->", n, "filas")

# ---------- 2025 ----------
wb = openpyxl.load_workbook("../privado/COSTOS POR UNIDAD 2025.xlsx", data_only=True, read_only=True)
for sheet in wb.sheetnames:
    if sheet in SKIP_2025:
        continue
    ws = wb[sheet]
    all_rows = list(ws.iter_rows(values_only=True))
    # buscar la fila de encabezado real (puede venir en la fila 0 o 1, ej. "GRÚA " trae una fila en blanco antes)
    header_row_i = None
    for i,r in enumerate(all_rows[:3]):
        if r and r[0] == "DIA":
            header_row_i = i
            break
    if header_row_i is None:
        print("2025", sheet, "-> SIN ENCABEZADO RECONOCIBLE, se omite")
        continue
    header = [ (h or "").strip() if isinstance(h,str) else h for h in all_rows[header_row_i] ]
    idx = {h:i for i,h in enumerate(header) if h}
    n = 0
    for row in all_rows[header_row_i+1:]:
        dia = row[idx.get("DIA",-1)] if idx.get("DIA",-1)>=0 else None
        if dia is None:
            continue
        if not isinstance(dia, (int, float)):
            # filas de "Total" u otros resúmenes al pie de la hoja, no son movimientos reales
            continue
        n += 1
        rows_out.append({
            "archivo_origen":"2025","hoja_origen":sheet,
            "dia": row[idx["DIA"]], "mes": row[idx["MES"]], "anio": row[idx["AÑO"]] if idx.get("AÑO",-1)>=0 else 2025,
            "factura": row[idx.get("FACTURA",-1)] if idx.get("FACTURA",-1)>=0 else None,
            "proveedor": row[idx.get("PROVEEDOR",-1)] if idx.get("PROVEEDOR",-1)>=0 else None,
            "concepto": row[idx.get("CONCEPTO",-1)] if idx.get("CONCEPTO",-1)>=0 else None,
            "equipo_col": row[idx["EQUIPO"]] if idx.get("EQUIPO",-1)>=0 else None,
            "empleado_col": row[idx["EMPLEADO"]] if idx.get("EMPLEADO",-1)>=0 else None,
            "cantidad": row[idx.get("CANTIDAD",-1)] if idx.get("CANTIDAD",-1)>=0 else None,
            "unidad": row[idx.get("UNIDAD",-1)] if idx.get("UNIDAD",-1)>=0 else None,
            "importe": to_num(row[idx.get("IMPORTE ",idx.get("IMPORTE",-1))]) if (idx.get("IMPORTE ",idx.get("IMPORTE",-1)))>=0 else None,
            "iva": to_num(row[idx.get("IVA",-1)]) if idx.get("IVA",-1)>=0 else None,
            "bruto": to_num(row[idx.get("BRUTO",-1)]) if idx.get("BRUTO",-1)>=0 else None,
            "neto": to_num(row[idx.get("NETO",-1)]) if idx.get("NETO",-1)>=0 else None,
            "tipo": row[idx.get("TIPO",-1)] if idx.get("TIPO",-1)>=0 else None,
        })
    print("2025", sheet, "->", n, "filas")

# ---------- 2026 ----------
wb = openpyxl.load_workbook("../privado/DASHBOARD MAQUINARIA.xlsx", data_only=True, read_only=True)
ws = wb["DASHBOARD"]
all_rows = list(ws.iter_rows(values_only=True))
header = None
header_i = None
for i,r in enumerate(all_rows):
    if r and "FECHA" in r and "MAQUINARIA" in r:
        header = r
        header_i = i
        break
idx = {h:j for j,h in enumerate(header) if h}
n = 0
for row in all_rows[header_i+1:]:
    fecha = row[idx.get("FECHA",-1)] if idx.get("FECHA",-1)>=0 else None
    if fecha is None:
        continue
    n += 1
    rows_out.append({
        "archivo_origen":"2026","hoja_origen":"REGISTRO",
        "dia": None, "mes": row[idx.get("MES",-1)], "anio": 2026, "fecha_completa": fecha,
        "factura": row[idx.get("FACTURA",-1)] if idx.get("FACTURA",-1)>=0 else None,
        "proveedor": row[idx.get("PROVEEDOR",-1)] if idx.get("PROVEEDOR",-1)>=0 else None,
        "concepto": row[idx.get("CONCEPTO",-1)] if idx.get("CONCEPTO",-1)>=0 else None,
        "equipo_col": row[idx.get("MAQUINARIA",-1)] if idx.get("MAQUINARIA",-1)>=0 else None,
        "empleado_col": None,
        "cantidad": row[idx.get("CANTIDAD",-1)] if idx.get("CANTIDAD",-1)>=0 else None,
        "unidad": row[idx.get("UNIDAD",-1)] if idx.get("UNIDAD",-1)>=0 else None,
        "importe": to_num(row[idx.get("P.U",-1)]) if idx.get("P.U",-1)>=0 else None,
        "iva": to_num(row[idx.get("IVA",-1)]) if idx.get("IVA",-1)>=0 else None,
        "bruto": to_num(row[idx.get("BRUTO",-1)]) if idx.get("BRUTO",-1)>=0 else None,
        "neto": to_num(row[idx.get("MONTO",-1)]) if idx.get("MONTO",-1)>=0 else None,
        "tipo": row[idx.get("TIPO",-1)] if idx.get("TIPO",-1)>=0 else None,
        "categoria_col": row[idx.get("CATEGORÍA",-1)] if idx.get("CATEGORÍA",-1)>=0 else None,
        "medio_pago": row[idx.get("MEDIO DE PAGO",-1)] if idx.get("MEDIO DE PAGO",-1)>=0 else None,
    })
print("2026 REGISTRO ->", n, "filas")

# ---------- guardar ----------
fields = sorted({k for r in rows_out for k in r.keys()})
with open("movimientos_crudo.csv","w",newline="",encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=fields)
    w.writeheader()
    for r in rows_out:
        w.writerow(r)

print("\nTOTAL:", len(rows_out), "filas crudas (esperado: 763+623+352 = 1738)")
por_anio = {}
for r in rows_out:
    por_anio[r["archivo_origen"]] = por_anio.get(r["archivo_origen"],0)+1
print(por_anio)
